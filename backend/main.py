
import re
from io import BytesIO

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pypdf import PdfReader


app = FastAPI(
    title="ResumeIQ API",
    description="Resume quality checks and job-role skill matching",
    version="4.0.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "https://resume-iq-nu-sooty.vercel.app",
    ],
    allow_credentials=True,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["*"],
)

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_PDF_PAGES = 50


# Canonical skill names and common aliases.
SKILL_ALIASES = {
    "Python": ["python"],
    "Java": ["java"],
    "JavaScript": ["javascript", "java script"],
    "TypeScript": ["typescript", "type script"],
    "React": ["react.js", "reactjs", "react"],
    "Node.js": ["node.js", "nodejs", "node js"],
    "Express.js": ["express.js", "expressjs", "express js"],
    "SQL": ["sql"],
    "MySQL": ["mysql", "my sql"],
    "PostgreSQL": ["postgresql", "postgres"],
    "MongoDB": ["mongodb", "mongo db"],
    "FastAPI": ["fastapi", "fast api"],
    "Django": ["django"],
    "Flask": ["flask"],
    "REST API": ["restful api", "rest apis", "rest api"],
    "GraphQL": ["graphql"],
    "HTML": ["html"],
    "CSS": ["css"],
    "Tailwind CSS": ["tailwind css", "tailwind"],
    "Bootstrap": ["bootstrap"],
    "AWS": ["amazon web services", "aws"],
    "Azure": ["microsoft azure", "azure"],
    "Google Cloud": ["google cloud", "gcp"],
    "Docker": ["docker"],
    "Kubernetes": ["kubernetes", "k8s"],
    "Jenkins": ["jenkins"],
    "GitHub Actions": ["github actions"],
    "GitLab CI/CD": ["gitlab ci/cd", "gitlab ci"],
    "CI/CD": [
        "continuous integration",
        "continuous delivery",
        "continuous deployment",
        "ci/cd",
    ],
    "Linux": ["linux"],
    "Bash": ["shell scripting", "bash"],
    "Git": ["git"],
    "GitHub": ["github"],
    "Terraform": ["terraform"],
    "Ansible": ["ansible"],
    "DevOps": ["devops", "dev ops"],
    "Python Automation": ["python automation"],
    "Machine Learning": ["machine learning", "ml"],
    "Deep Learning": ["deep learning"],
    "Natural Language Processing": [
        "natural language processing",
        "nlp",
    ],
    "Data Analysis": ["data analysis", "data analytics"],
    "Pandas": ["pandas"],
    "NumPy": ["numpy"],
    "Power BI": ["power bi"],
    "Tableau": ["tableau"],
    "Excel": ["microsoft excel", "excel"],
    "n8n": ["n8n"],
    "OpenAI API": ["openai api"],
    "Figma": ["figma"],
    "Agile": ["agile"],
    "Communication": ["communication"],
}


# Required skills for supported roles.
ROLE_SKILLS = {
    "python developer": [
        "Python", "FastAPI", "Django", "SQL", "Git",
    ],
    "python": [
        "Python", "FastAPI", "Django", "SQL", "Git",
    ],
    "backend developer": [
        "Python", "FastAPI", "SQL", "REST API", "Git",
    ],
    "devops engineer": [
        "Linux", "Docker", "Kubernetes", "Jenkins",
        "AWS", "CI/CD", "Git",
    ],
    "devops": [
        "Linux", "Docker", "Kubernetes", "Jenkins",
        "AWS", "CI/CD", "Git",
    ],
    "cloud engineer": [
        "Linux", "AWS", "Azure", "Docker", "Terraform", "Git",
    ],
    "frontend developer": [
        "HTML", "CSS", "JavaScript", "TypeScript", "React",
    ],
    "frontend": [
        "HTML", "CSS", "JavaScript", "TypeScript", "React",
    ],
    "full stack developer": [
        "HTML", "CSS", "JavaScript", "React", "Node.js", "SQL",
    ],
    "data analyst": [
        "SQL", "Python", "Power BI", "Data Analysis", "Excel",
    ],
    "machine learning engineer": [
        "Python", "Machine Learning", "Deep Learning", "SQL",
    ],
    "machine learning": [
        "Python", "Machine Learning", "Deep Learning", "SQL",
    ],
    "ai automation engineer": [
        "Python", "n8n", "REST API", "SQL", "Git",
    ],
    "automation engineer": [
        "Python", "REST API", "Git", "Linux",
    ],
}


ROLE_ALIASES = {
    "front end developer": "frontend developer",
    "front-end developer": "frontend developer",
    "fullstack developer": "full stack developer",
    "back end developer": "backend developer",
    "back-end developer": "backend developer",
    "ai automation": "ai automation engineer",
}


SECTION_PATTERNS = {
    "Education": [
        r"education",
        r"academic background",
        r"educational qualifications?",
        r"qualifications?",
    ],
    "Skills": [
        r"skills",
        r"technical skills",
        r"core competencies",
        r"key skills",
        r"technologies",
    ],
    "Projects": [
        r"projects",
        r"personal projects",
        r"academic projects",
        r"project experience",
    ],
    "Experience": [
        r"experience",
        r"work experience",
        r"professional experience",
        r"internships?",
        r"employment history",
    ],
    "Certifications": [
        r"certifications?",
        r"certificates",
        r"courses",
        r"licenses",
    ],
}


def normalize_text(text: str) -> str:
    """Normalize whitespace and case."""
    return re.sub(r"\s+", " ", text).strip().lower()


def contains_term(text: str, term: str) -> bool:
    """Match whole terms instead of unrelated substrings."""
    pattern = r"(?<!\w)" + re.escape(term) + r"(?!\w)"
    return re.search(pattern, text, re.IGNORECASE) is not None


def extract_text(filename: str, content: bytes) -> str:
    """Extract text from text-based PDF or DOCX files."""
    extension = filename.lower().rsplit(".", 1)[-1]

    try:
        if extension == "pdf":
            reader = PdfReader(BytesIO(content), strict=False)

            if reader.is_encrypted:
                try:
                    if not reader.decrypt(""):
                        raise HTTPException(
                            status_code=400,
                            detail=(
                                "This PDF is password-protected. "
                                "Upload an unlocked PDF."
                            ),
                        )
                except HTTPException:
                    raise
                except Exception as exc:
                    raise HTTPException(
                        status_code=400,
                        detail="Unable to open the protected PDF.",
                    ) from exc

            if len(reader.pages) > MAX_PDF_PAGES:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"PDF contains too many pages. "
                        f"Maximum supported: {MAX_PDF_PAGES}."
                    ),
                )

            return "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        if extension == "docx":
            document = Document(BytesIO(content))

            parts = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            # Include content stored inside Word tables.
            for table in document.tables:
                for row in table.rows:
                    parts.append(
                        " ".join(cell.text for cell in row.cells)
                    )

            return "\n".join(parts)

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=(
                "Could not read this resume. "
                "Upload a valid text-based PDF or DOCX."
            ),
        ) from exc

    raise HTTPException(
        status_code=400,
        detail="Unsupported file type. Upload a PDF or DOCX resume.",
    )


def detect_skills(text: str) -> list[str]:
    """Detect unique skills using canonical names."""
    normalized = normalize_text(text)
    found = []

    for canonical_name, aliases in SKILL_ALIASES.items():
        if any(
            contains_term(normalized, alias)
            for alias in aliases
        ):
            found.append(canonical_name)

    return found


def detect_sections(text: str) -> list[str]:
    """Detect common section headings line by line."""
    lines = [
        re.sub(r"\s+", " ", line).strip().lower()
        for line in text.splitlines()
        if line.strip()
    ]

    found = []

    for section, patterns in SECTION_PATTERNS.items():
        for line in lines:
            cleaned = re.sub(r"^[\W_]+|[\W_]+$", "", line)

            if len(cleaned) > 90:
                continue

            if any(
                re.fullmatch(
                    pattern + r"\s*[:\-–—]?",
                    cleaned,
                    re.IGNORECASE,
                )
                for pattern in patterns
            ):
                found.append(section)
                break

    return found


def analyze_text(text: str, job_role: str) -> dict:
    """Calculate a transparent resume-quality estimate."""

    # 1. Basic information
    word_count = len(re.findall(r"\b[\w+#./-]+\b", text))
    normalized = normalize_text(text)

    email_found = bool(
        re.search(
            r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",
            text,
            re.IGNORECASE,
        )
    )

    phone_found = any(
        8 <= len(re.sub(r"\D", "", match)) <= 15
        for match in re.findall(
            r"(?<!\w)\+?\d[\d\s().-]{6,}\d(?!\w)",
            text,
        )
    )

    date_found = bool(
        re.search(
            r"\b(?:19|20)\d{2}\b|"
            r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
            r"[a-z]*\.?\s+\d{4}\b",
            text,
            re.IGNORECASE,
        )
    )

    quantified_achievement = bool(
        re.search(
            r"\b\d+(?:\.\d+)?\s*"
            r"(?:%|percent|users|customers|projects|hours|days|"
            r"weeks|months|years|requests|seconds|minutes)\b",
            text,
            re.IGNORECASE,
        )
    )

    # 2. Skills and sections
    found_skills = detect_skills(text)
    sections_found = detect_sections(text)

    all_sections = [
        "Education",
        "Skills",
        "Projects",
        "Experience",
        "Certifications",
    ]

    missing_sections = [
        section
        for section in all_sections
        if section not in sections_found
    ]

    # 3. Role normalization
    normalized_role = normalize_text(job_role)
    role_key = ROLE_ALIASES.get(
        normalized_role,
        normalized_role,
    )

    expected_skills = ROLE_SKILLS.get(role_key, [])

    matched_role_skills = [
        skill
        for skill in expected_skills
        if skill in found_skills
    ]

    missing_role_skills = [
        skill
        for skill in expected_skills
        if skill not in matched_role_skills
    ]

    if expected_skills:
        skill_match_score = round(
            len(matched_role_skills) / len(expected_skills) * 100
        )
        skill_quality_points = round(
            skill_match_score * 0.20
        )
    else:
        # Unsupported roles must not receive a fabricated role-match score.
        skill_match_score = None
        skill_quality_points = min(len(found_skills) * 2, 20)

    # 4. Look for evidence of project and experience details.
    project_action_found = bool(
        re.search(
            r"\b(built|developed|created|deployed|implemented|"
            r"designed|automated|configured|integrated|maintained)\b",
            normalized,
        )
    )

    experience_action_found = bool(
        re.search(
            r"\b(worked|managed|delivered|supported|resolved|"
            r"collaborated|optimized|implemented|deployed|maintained)\b",
            normalized,
        )
    )

    project_points = (
        10
        if "Projects" in sections_found and project_action_found
        else 5
        if "Projects" in sections_found
        else 0
    )

    experience_points = (
        10
        if "Experience" in sections_found and experience_action_found
        else 5
        if "Experience" in sections_found
        else 0
    )

    # 5. Resume quality score (maximum 100 points).
    # This is a heuristic, not a standardized ATS score.
    section_points = min(len(sections_found) * 5, 25)

    contact_points = (
        (5 if email_found else 0)
        + (5 if phone_found else 0)
    )

    if 150 <= word_count <= 800:
        length_points = 10
    elif 100 <= word_count < 150 or 801 <= word_count <= 1000:
        length_points = 6
    elif word_count > 0:
        length_points = 2
    else:
        length_points = 0

    date_points = 5 if date_found else 0
    achievement_points = 10 if quantified_achievement else 0

    score_breakdown = {
        "sections": section_points,  # max 25
        "skills": skill_quality_points,  # max 20
        "email": 5 if email_found else 0,
        "phone": 5 if phone_found else 0,
        "length": length_points,  # max 10
        "dates": date_points,  # max 5
        "measurable_achievements": achievement_points,  # max 10
        "project_details": project_points,  # max 10
        "experience_details": experience_points,  # max 10
    }

    quality_score = min(sum(score_breakdown.values()), 100)

    # 6. Strengths and improvements
    strengths = []
    improvements = []

    if email_found:
        strengths.append("An email address was detected.")
    else:
        improvements.append("Add a professional email address.")

    if phone_found:
        strengths.append("A possible contact phone number was detected.")
    else:
        improvements.append("Consider adding a contact phone number.")

    if "Education" in sections_found:
        strengths.append("An Education section was detected.")
    else:
        improvements.append("Add a clearly labelled Education section.")

    if "Skills" in sections_found and found_skills:
        strengths.append(
            "A Skills section and technical skills were detected."
        )
    else:
        improvements.append(
            "Add a clearly labelled Skills section with relevant skills."
        )

    if "Projects" in sections_found:
        strengths.append("A Projects section was detected.")
        if not project_action_found:
            improvements.append(
                "Explain what you built, the tools used, and your contribution."
            )
    else:
        improvements.append(
            "Add relevant projects and explain your contribution."
        )

    if "Experience" in sections_found:
        strengths.append(
            "An Experience or internship section was detected."
        )
        if not experience_action_found:
            improvements.append(
                "Explain your responsibilities and practical work."
            )
    else:
        improvements.append(
            "If applicable, include internships or practical experience."
        )

    if "Certifications" not in sections_found:
        improvements.append(
            "Add relevant certifications or courses if you have them."
        )

    if not date_found:
        improvements.append(
            "Include dates for education, projects, courses, or experience."
        )

    if not quantified_achievement:
        improvements.append(
            "Where accurate, add measurable project outcomes or results."
        )

    if expected_skills and missing_role_skills:
        improvements.append(
            "Consider demonstrating these target-role skills if you know them: "
            + ", ".join(missing_role_skills)
            + "."
        )

    if word_count < 100:
        improvements.append(
            "Very little text was extracted. Check PDF text extraction "
            "and whether all pages were read."
        )
    elif word_count > 1000:
        improvements.append(
            "Consider removing repetitive content and keeping the resume concise."
        )

    if not improvements:
        improvements.append(
            "Keep your resume accurate and update it with relevant evidence."
        )

    # 7. Preserve the response fields expected by the existing frontend.
    return {
        "score": quality_score,
        "quality_score": quality_score,
        "skill_match_score": skill_match_score,
        "score_type": "rule-based estimate",
        "word_count": word_count,
        "skills": found_skills,
        "sections_found": sections_found,
        "missing_sections": missing_sections,
        "strengths": strengths,
        "improvements": improvements,
        "job_role": job_role.strip() or "Not specified",
        "matched_role_skills": matched_role_skills,
        "missing_role_skills": missing_role_skills,
        "note": (
            "Scores are rule-based guidance, not an official ATS score "
            "or a prediction of hiring success. Skill matching checks "
            "only the configured skills for the selected role."
        ),
        "score_breakdown": score_breakdown,
        "email_detected": email_found,
        "phone_detected": phone_found,
        "date_detected": date_found,
        "measurable_achievement_detected": quantified_achievement,
        "detected_sections_count": len(sections_found),
    }


@app.get("/")
def home():
    return {
        "message": "ResumeIQ API is running",
        "version": app.version,
        "docs": "/docs",
    }


@app.get("/health")
def health():
    return {
        "status": "ok",
        "service": "ResumeIQ API",
    }


@app.post("/analyze")
async def analyze_resume(
    file: UploadFile = File(...),
    job_role: str = Form(""),
):
    filename = file.filename or ""

    if not filename:
        raise HTTPException(
            status_code=400,
            detail="Please select a resume file.",
        )

    if not filename.lower().endswith((".pdf", ".docx")):
        raise HTTPException(
            status_code=400,
            detail="Please upload a PDF or DOCX resume.",
        )

    try:
        # Read only the maximum allowed size plus one byte.
        content = await file.read(MAX_FILE_SIZE + 1)
    finally:
        await file.close()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=400,
            detail="File is too large. Maximum size is 5 MB.",
        )

    text = extract_text(filename, content)

    if not text.strip():
        raise HTTPException(
            status_code=400,
            detail=(
                "No readable text was found. Try a text-based PDF or DOCX. "
                "Scanned image-only PDFs are not supported."
            ),
        )

    return analyze_text(text, job_role)
